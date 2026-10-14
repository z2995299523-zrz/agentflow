"""元数据标准模板的测试"""
from openpyxl import load_workbook

from metadata.template import GUIDE_SHEET_NAME, SHEETS, main, write_blank_template


def test_four_sheets_with_expected_names():
    assert [sheet.name for sheet in SHEETS] == ["表清单", "字段清单", "血缘清单", "关联关系"]
    assert [sheet.key for sheet in SHEETS] == ["tables", "columns", "lineage", "relations"]


def test_only_relation_sheet_is_optional():
    assert [sheet.optional for sheet in SHEETS] == [False, False, False, True]


def test_required_headers():
    required = {sheet.name: [c.header for c in sheet.columns if c.required] for sheet in SHEETS}
    assert required == {
        "表清单": ["层级", "模式名", "表英文名", "表中文名", "表类型"],
        "字段清单": ["表名", "字段英文名", "字段中文名"],
        "血缘清单": ["目标表", "目标字段", "源表", "源字段"],
        "关联关系": ["左表", "左字段", "右表", "右字段", "对应关系"],
    }


def test_field_names_are_unique_within_sheet():
    for sheet in SHEETS:
        fields = [column.field for column in sheet.columns]
        assert len(fields) == len(set(fields)), sheet.name


def test_write_blank_template(tmp_path):
    path = write_blank_template(tmp_path / "out" / "模板.xlsx")

    workbook = load_workbook(path)
    assert workbook.sheetnames == [GUIDE_SHEET_NAME, "表清单", "字段清单", "血缘清单", "关联关系"]
    for sheet in SHEETS:
        worksheet = workbook[sheet.name]
        assert [cell.value for cell in worksheet[1]] == [c.header for c in sheet.columns]
        assert worksheet.max_row == 1          # 只有表头，没有数据
    guide_rows = list(workbook[GUIDE_SHEET_NAME].iter_rows(values_only=True))
    assert guide_rows[0] == ("清单", "列名", "是否必填", "说明")
    assert len(guide_rows) == 1 + sum(len(sheet.columns) for sheet in SHEETS)
    assert ("关联关系（可选）", "对应关系", "是", "只能填：N:1、1:1、1:N") in guide_rows


def test_main_writes_to_given_path(tmp_path, capsys):
    target = tmp_path / "t.xlsx"

    assert main([str(target)]) == 0
    assert target.exists()
    assert "空白模板已生成" in capsys.readouterr().out
