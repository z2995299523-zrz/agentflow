"""资料读取的测试"""
import pytest
from openpyxl import Workbook

from metadata.loader import load_raw
from metadata.template import SHEETS, WORKBOOK_NAME
from tests.helpers import write_mart


def test_load_csv_mart(tiny_mart):
    raw = load_raw(tiny_mart)

    assert raw.missing_sheets == []
    assert raw.missing_columns == {}
    assert {key: len(rows) for key, rows in raw.rows.items()} == {
        "tables": 3, "columns": 9, "lineage": 3, "relations": 1,
    }
    first = raw.rows["tables"][0]
    assert (first.sheet, first.row_no) == ("表清单", 2)
    assert first.values["table_cn"] == "账户表"
    assert first.values["grain_columns"] == "data_dt,acct_no"


def test_identifiers_are_lowercased_but_text_is_kept(tmp_path, tiny_sheets):
    tiny_sheets["表清单"][0][1:3] = ["ODS", "ACCT"]
    tiny_sheets["表清单"][0][7] = "Daily Snapshot"

    values = load_raw(write_mart(tmp_path / "m", tiny_sheets)).rows["tables"][0].values

    assert (values["schema_name"], values["table_name"]) == ("ods", "acct")
    assert values["description"] == "Daily Snapshot"


def test_cells_are_stripped_and_blank_rows_skipped(tmp_path, tiny_sheets):
    tiny_sheets["字段清单"][0][2] = "  数据日期  "
    tiny_sheets["字段清单"].insert(1, ["", "", "", "", "", "", ""])

    rows = load_raw(write_mart(tmp_path / "m", tiny_sheets)).rows["columns"]

    assert len(rows) == 9
    assert rows[0].values["column_cn"] == "数据日期"
    assert rows[1].row_no == 4      # 空行占了第 3 行，行号仍对应文件里的真实位置


def test_optional_sheet_may_be_absent(tmp_path, tiny_sheets):
    del tiny_sheets["关联关系"]

    raw = load_raw(write_mart(tmp_path / "m", tiny_sheets))

    assert raw.missing_sheets == []
    assert raw.rows["relations"] == []


def test_required_sheet_missing(tmp_path, tiny_sheets):
    del tiny_sheets["血缘清单"]

    raw = load_raw(write_mart(tmp_path / "m", tiny_sheets))

    assert raw.missing_sheets == ["血缘清单"]
    assert raw.rows["lineage"] == []


def test_required_column_missing(tmp_path, tiny_sheets):
    mart_dir = write_mart(tmp_path / "m", tiny_sheets)
    (mart_dir / "字段清单.csv").write_text(
        "表名,字段英文名,数据类型\nods.acct,bal,numeric\n", encoding="utf-8")

    assert load_raw(mart_dir).missing_columns == {"字段清单": ["字段中文名"]}


def test_optional_column_may_be_absent(tmp_path, tiny_sheets):
    mart_dir = write_mart(tmp_path / "m", tiny_sheets)
    # 血缘清单不带「加工逻辑」列，与真实的血缘清单一致
    (mart_dir / "血缘清单.csv").write_text(
        "目标表,目标字段,源表,源字段,加工脚本\nads.cust_bal,bal,ods.acct,bal,load.sql\n",
        encoding="utf-8-sig")

    raw = load_raw(mart_dir)

    assert raw.missing_columns == {}
    assert raw.rows["lineage"][0].values == {
        "target_table": "ads.cust_bal", "target_column": "bal",
        "source_table": "ods.acct", "source_column": "bal",
        "transform": "", "script_name": "load.sql",
    }


def test_workbook_takes_precedence_over_csv(tmp_path, tiny_sheets):
    mart_dir = write_mart(tmp_path / "m", tiny_sheets)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for spec in SHEETS:
        workbook.create_sheet(spec.name).append([column.header for column in spec.columns])
    workbook["表清单"].append(["贴源层", "ODS", "only_in_xlsx", "工作簿里的表", "维表", "id", None, None])
    workbook.save(mart_dir / WORKBOOK_NAME)

    rows = load_raw(mart_dir).rows["tables"]

    assert [row.values["table_name"] for row in rows] == ["only_in_xlsx"]
    assert rows[0].values["schema_name"] == "ods"
    assert rows[0].values["date_column"] == ""      # 空单元格读成空文本


def test_missing_directory_raises(tmp_path):
    with pytest.raises(FileNotFoundError, match="资料目录不存在"):
        load_raw(tmp_path / "nope")


def test_csv_saved_by_excel_in_gbk_is_read(tmp_path, tiny_sheets):
    mart_dir = write_mart(tmp_path / "m", tiny_sheets)
    path = mart_dir / "表清单.csv"
    path.write_bytes(path.read_text(encoding="utf-8-sig").encode("gbk"))    # 模拟 Excel 另存为 CSV

    raw = load_raw(mart_dir)

    assert raw.missing_columns == {}
    assert raw.rows["tables"][0].values["table_cn"] == "账户表"


def test_csv_in_unknown_encoding_gives_readable_error(tmp_path, tiny_sheets):
    mart_dir = write_mart(tmp_path / "m", tiny_sheets)
    (mart_dir / "表清单.csv").write_bytes(bytes([0xFF, 0xFE, 0xFF]))

    with pytest.raises(ValueError, match="无法识别 表清单.csv 的编码"):
        load_raw(mart_dir)


def test_spaces_around_headers_and_sheet_names_are_ignored(tmp_path, tiny_sheets):
    mart_dir = write_mart(tmp_path / "m", tiny_sheets)
    (mart_dir / "字段清单.csv").write_text(
        " 表名 ,字段英文名 , 字段中文名,数据类型\nods.acct,bal,余额,numeric\n", encoding="utf-8")
    assert load_raw(mart_dir).missing_columns == {}

    workbook = Workbook()
    workbook.remove(workbook.active)
    for spec in SHEETS:
        workbook.create_sheet(spec.name + " ").append([c.header + " " for c in spec.columns])
    workbook.save(mart_dir / WORKBOOK_NAME)
    raw = load_raw(mart_dir)
    assert raw.missing_sheets == [] and raw.missing_columns == {}
