"""
资料读取
把资料目录里的四张清单读成原始行。只负责读取和统一大小写，不做业务校验
"""
import csv
import io
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from metadata.template import SHEETS, WORKBOOK_NAME, SheetSpec

# 依次尝试的 CSV 编码。中文 Windows 上用 Excel 另存 CSV 时默认是 GBK 而不是 UTF-8
CSV_ENCODINGS = ("utf-8-sig", "gb18030")


@dataclass
class RawRow:
    """清单里的一行原始记录"""

    sheet: str                 # 清单名
    row_no: int                # 行号，表头是第 1 行
    values: dict[str, str]     # 内部字段名 → 单元格文本


@dataclass
class RawMart:
    """一套资料的原始内容"""

    mart_dir: Path
    rows: dict[str, list[RawRow]] = field(default_factory=dict)            # 清单内部键 → 行
    missing_sheets: list[str] = field(default_factory=list)                # 缺失的必需清单名
    missing_columns: dict[str, list[str]] = field(default_factory=dict)    # 清单名 → 缺失的必填列名


def _cell_text(value: object) -> str:
    """把单元格的值统一成去掉首尾空格的文本"""
    if value is None:
        return ""
    return str(value).strip()


def _read_xlsx(path: Path) -> dict[str, list[list[str]]]:
    """读取工作簿，返回 清单名 → 行（含表头）"""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return {
            worksheet.title.strip(): [
                [_cell_text(cell) for cell in row]
                for row in worksheet.iter_rows(values_only=True)
            ]
            for worksheet in workbook.worksheets
        }
    finally:
        workbook.close()


def _read_csv(path: Path) -> list[list[str]]:
    """读取 CSV 文件，返回行（含表头）

    先按 UTF-8 解码（带不带 BOM 都可以），失败再按 GBK 解码。

    Raises:
        ValueError: 两种编码都无法解码
    """
    data = path.read_bytes()
    for encoding in CSV_ENCODINGS:
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"无法识别 {path.name} 的编码，请另存为 UTF-8 或 GBK 编码的 CSV")
    return [[_cell_text(cell) for cell in row]
            for row in csv.reader(io.StringIO(text, newline=""))]


def _to_rows(spec: SheetSpec, grid: list[list[str]], raw: RawMart) -> list[RawRow]:
    """按清单定义把表格转成原始行，同时登记缺失的必填列"""
    header = grid[0] if grid else []
    positions = {name: index for index, name in enumerate(header) if name}
    missing = [column.header for column in spec.columns
               if column.required and column.header not in positions]
    if missing:
        raw.missing_columns[spec.name] = missing

    rows: list[RawRow] = []
    for row_no, cells in enumerate(grid[1:], start=2):
        if not any(cells):
            continue  # 跳过空行
        values: dict[str, str] = {}
        for column in spec.columns:
            index = positions.get(column.header)
            text = cells[index] if index is not None and index < len(cells) else ""
            values[column.field] = text.lower() if column.identifier else text
        rows.append(RawRow(sheet=spec.name, row_no=row_no, values=values))
    return rows


def load_raw(mart_dir: Path) -> RawMart:
    """读取一套资料

    目录里有 metadata.xlsx 时读工作簿的四个 Sheet，否则读四个同名 CSV 文件。
    可选的清单或可选的列不存在时按空处理。

    Args:
        mart_dir: 资料目录

    Returns:
        原始资料，尚未校验

    Raises:
        FileNotFoundError: 资料目录不存在
    """
    mart_dir = Path(mart_dir)
    if not mart_dir.is_dir():
        raise FileNotFoundError(f"资料目录不存在: {mart_dir}")

    raw = RawMart(mart_dir=mart_dir)
    workbook_path = mart_dir / WORKBOOK_NAME
    sheets = _read_xlsx(workbook_path) if workbook_path.exists() else None

    for spec in SHEETS:
        if sheets is not None:
            grid = sheets.get(spec.name)
        else:
            csv_path = mart_dir / f"{spec.name}.csv"
            grid = _read_csv(csv_path) if csv_path.exists() else None

        if grid is None:
            if not spec.optional:
                raw.missing_sheets.append(spec.name)
            raw.rows[spec.key] = []
            continue
        raw.rows[spec.key] = _to_rows(spec, grid, raw)
    return raw
